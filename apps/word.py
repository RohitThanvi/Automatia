"""
Word document automation (spec section 11).

Same design rationale as apps/powerpoint.py: author the actual
document content directly via python-docx rather than driving Word's
UI keystroke by keystroke. This makes "write a 2000-word essay and
save it as a Word document" a single deterministic tool call instead
of thousands of simulated keypresses.
"""

from __future__ import annotations

from pathlib import Path

from core.config import get_config
from core.logging_setup import get_logger
from core.tool_registry import registry

log = get_logger()


def _resolve_output_path(filename: str) -> Path:
    if not filename.lower().endswith(".docx"):
        filename += ".docx"
    p = Path(filename).expanduser()
    if not p.is_absolute():
        p = get_config().resolved_workspace_dir() / p
    p.parent.mkdir(parents=True, exist_ok=True)
    return p


@registry.register(
    "create_document",
    (
        "Create a formatted Word (.docx) document from a title and a list of sections, each "
        "with a heading and body paragraph(s). Generate the actual written content first "
        "(this tool only formats and saves it, it does not write the content for you)."
    ),
    {
        "title": {"type": "string", "description": "Document title, shown as the top heading."},
        "output_filename": {"type": "string", "description": "Filename to save as, e.g. 'Indus_Valley_Essay.docx'."},
        "sections": {
            "type": "array",
            "description": "Document sections, in order.",
            "items": {
                "type": "object",
                "properties": {
                    "heading": {"type": "string", "description": "Optional section heading; omit for a single flowing essay."},
                    "paragraphs": {"type": "array", "items": {"type": "string"}},
                },
            },
        },
        "open_after": {"type": "boolean", "description": "Open the file in Word after saving. Default true."},
    },
    required=["title", "output_filename", "sections"],
)
def create_document(
    title: str,
    output_filename: str,
    sections: list[dict],
    open_after: bool = True,
) -> dict:
    from docx import Document
    from docx.shared import Pt

    doc = Document()

    title_heading = doc.add_heading(title, level=0)

    for section in sections:
        heading = section.get("heading")
        if heading:
            doc.add_heading(heading, level=1)
        for para_text in section.get("paragraphs", []):
            p = doc.add_paragraph(str(para_text))
            p.style.font.size = Pt(11)

    output_path = _resolve_output_path(output_filename)
    doc.save(str(output_path))

    word_count = sum(
        len(str(p).split())
        for s in sections
        for p in s.get("paragraphs", [])
    )
    log.info(f"Saved document: {output_path} (~{word_count} words)")

    result = {"ok": True, "path": str(output_path), "word_count": word_count}

    if open_after:
        from computer.windows import open_application

        open_result = open_application(str(output_path))
        result["opened"] = open_result.get("ok", False)

    return result
