import pytest


@pytest.fixture(autouse=True)
def temp_workspace(tmp_path, monkeypatch):
    from core.config import reload_config

    cfg = reload_config()
    cfg.paths.workspace_dir = str(tmp_path)
    yield tmp_path
    reload_config()


def test_create_presentation_produces_real_pptx(temp_workspace, monkeypatch):
    from apps.powerpoint import create_presentation

    # Don't actually try to launch PowerPoint during tests.
    monkeypatch.setattr(
        "computer.windows.open_application", lambda path: {"ok": True}
    )

    result = create_presentation(
        title="Indus Valley Civilization",
        output_filename="test_deck.pptx",
        slides=[
            {"heading": "Overview", "bullets": ["Bronze Age civilization", "Located in South Asia"]},
            {"heading": "Cities", "bullets": ["Harappa", "Mohenjo-daro"]},
        ],
        open_after=True,
    )

    assert result["ok"] is True
    assert result["slide_count"] == 3  # title slide + 2 content slides
    import os

    assert os.path.exists(result["path"])

    # Verify the file is a real, parseable pptx with the expected content.
    from pptx import Presentation

    prs = Presentation(result["path"])
    assert len(prs.slides) == 3
    assert prs.slides[0].shapes.title.text == "Indus Valley Civilization"
    assert prs.slides[1].shapes.title.text == "Overview"


def test_create_presentation_without_opening(temp_workspace):
    from apps.powerpoint import create_presentation

    result = create_presentation(
        title="Test",
        output_filename="no_open.pptx",
        slides=[{"heading": "A", "bullets": ["b1"]}],
        open_after=False,
    )
    assert result["ok"] is True
    assert "opened" not in result
